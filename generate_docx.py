"""Generate DESIGN.docx from DESIGN.md content."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)


def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # Set shading
    shading = run._element.get_or_add_rPr()
    shdng = shading.makeelement(
        qn("w:shd"),
        {
            qn("w:val"): "clear",
            qn("w:fill"): "F0F0F0",
        },
    )
    shading.append(shdng)


def build_docx():
    doc = Document()

    # Styles
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    for level in range(1, 4):
        hs = doc.styles[f"Heading {level}"]
        hs.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    # Title
    title = doc.add_heading("Design Document", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- Architecture Overview ---
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph(
        "The agent uses a hybrid FSM + Claude API architecture. A deterministic "
        "state machine controls conversation flow, verification gates, and PII "
        "isolation, while Claude handles natural language understanding (entity "
        "extraction via tool_use) and natural language generation (conversational responses)."
    )

    add_code_block(
        doc,
        "                     +----------------------------+\n"
        "                     |          Agent             |\n"
        "                     |   (Hybrid FSM + Claude)    |\n"
        "                     +-------------+--------------+\n"
        "          +---------------+--------+--------+---------------+\n"
        "          |               |        |        |               |\n"
        "   +------+------+ +-----+-----+  |  +-----+------+ +-----+------+\n"
        "   | LLM Client  | |  Memory   |  |  |Verification| | Validators |\n"
        "   | (Claude API)| | (3-tier)  |  |  |  Service   | |(Card, Date)|\n"
        "   +------+------+ +-----------+  |  +------------+ +------------+\n"
        "          |                        |\n"
        "   +------+------+         +------+------+\n"
        "   | Tool Defs   |         | API Client  |\n"
        "   | + Prompts   |         | (Abstract)  |\n"
        "   | + PIIScanner|         +------+------+\n"
        "   +-------------+                |\n"
        "                          +-------+------+\n"
        "                          | httpx impl   |\n"
        "                          +--------------+",
    )

    doc.add_heading("Why Hybrid, Not Full-LLM or Full-Rule-Based", level=2)
    items = [
        (
            "Full-LLM risk:",
            "An unconstrained LLM could skip verification steps, "
            "leak PII via prompt injection, or produce non-deterministic verification "
            "behavior. These are unacceptable in a payment flow.",
        ),
        (
            "Full-rule-based limitation:",
            "Regex parsing is brittle -- it misses "
            "phrasing variations and can't handle out-of-order input naturally. "
            "An LLM excels at understanding diverse input.",
        ),
        (
            "Hybrid advantage:",
            "The state machine enforces invariants (verification "
            'gate, step ordering, retry limits), while Claude handles the "messy" '
            'parts (parsing "my birthday is May 14th 1990" or "I want to pay the full amount").',
        ),
    ]
    for label, text in items:
        p = doc.add_paragraph()
        run = p.add_run(label + " ")
        run.bold = True
        p.add_run(text)

    doc.add_heading("State Machine", level=2)
    add_code_block(
        doc,
        "GREETING --> AWAITING_ACCOUNT_ID --> AWAITING_NAME --> AWAITING_SECONDARY\n"
        "                                                              |\n"
        "                                          +-------------------+\n"
        "                                          v\n"
        "                                   BALANCE_DISCLOSED --> AWAITING_AMOUNT\n"
        "                                                              |\n"
        "                                          +-------------------+\n"
        "                                          v\n"
        "                                    COLLECTING_CARD --> PAYMENT_COMPLETE --> CLOSED",
    )
    doc.add_paragraph(
        "Any verification or payment failure beyond the retry limit transitions directly to CLOSED."
    )

    # --- Memory Architecture ---
    doc.add_heading("Memory Architecture (3-Tier)", level=1)
    doc.add_paragraph(
        "Inspired by mem0's extract-consolidate-retrieve pipeline and "
        "LangGraph's reducer-driven state:"
    )
    add_table(
        doc,
        ["Tier", "Class", "Purpose", "Mutability"],
        [
            [
                "1",
                "WorkingMemory",
                "Structured state -- source of truth for FSM",
                "Read/write by state machine",
            ],
            [
                "2",
                "ConversationMemory",
                "Sliding window + summary for Claude context",
                "Append-only messages",
            ],
            [
                "3",
                "SemanticMemory",
                "Extracted facts injected into system prompt",
                "Append-only facts",
            ],
        ],
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("EntityBuffer")
    run.bold = True
    p.add_run(
        " (part of WorkingMemory) handles slot-filling for out-of-order input. "
        'When a user says "Hi, I\'m Nithin Jain, my account is ACC1001", both the '
        "name and account_id are buffered. The state machine drains entities in "
        "flow order -- account_id first, then name -- so no information is lost."
    )

    # --- Security Guardrails ---
    doc.add_heading("Security Guardrails", level=1)

    doc.add_heading("PII Safety (Defense-in-Depth)", level=2)
    doc.add_paragraph("Four layers of protection:")
    pii_items = [
        (
            "Architecture-level isolation:",
            "DOB, Aadhaar last 4, and pincode are "
            "NEVER sent to Claude. The state machine performs verification locally "
            "using VerificationService. Claude only sees sanitized status messages "
            '("Entities extracted.").',
        ),
        (
            "System prompt rules:",
            "Every state-specific prompt includes hard rules: "
            '"NEVER reveal the user\'s date of birth, Aadhaar number, or pincode."',
        ),
        (
            "Post-response PIIScanner:",
            "Before any response reaches the user, "
            "PIIScanner checks for leaked sensitive values using word-boundary regex "
            "matching with date format variants (YYYY-MM-DD, DD/MM/YYYY, MM/DD/YYYY). "
            "If found, the entire response is replaced with a safe fallback.",
        ),
        (
            "Model repr redaction:",
            "AccountData.__repr__() and CardDetails.__repr__() "
            "mask sensitive fields, preventing accidental PII exposure in logs, "
            "tracebacks, or debug output.",
        ),
    ]
    for i, (label, text) in enumerate(pii_items, 1):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(label + " ")
        run.bold = True
        p.add_run(text)

    doc.add_heading("Card Data Protection", level=2)
    card_items = [
        (
            "Conversation memory redaction:",
            "Card numbers and CVVs are masked "
            "(_redact_card_data()) before being stored in conversation memory, so "
            "they never reach Claude's context on subsequent turns.",
        ),
        (
            "Tool call sanitization:",
            "When Claude extracts card details via "
            "extract_card_details, the tool input is sanitized (card number masked "
            "to ****XXXX, CVV replaced with ***) before storage.",
        ),
        (
            "Immediate cleanup:",
            "Card details are cleared from working memory on "
            "payment completion, failure, or session close via _close_session().",
        ),
        (
            "HTTPS enforcement:",
            "PaymentAPIClient rejects non-HTTPS base URLs, "
            "preventing plaintext transmission of card data.",
        ),
    ]
    for label, text in card_items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(label + " ")
        run.bold = True
        p.add_run(text)

    doc.add_heading("Input Validation Guardrails", level=2)
    val_items = [
        (
            "Entity format validation:",
            "EntityBuffer.fill() validates all extracted "
            "entities before accepting them -- account IDs must match ACC\\d+, names "
            "must be >= 2 characters, Aadhaar must be exactly 4 digits, pincode "
            "exactly 6 digits. This prevents hallucinated extractions from "
            'tool_choice: "any" from triggering spurious state transitions.',
        ),
        (
            "Input length limit:",
            "User input is truncated to 500 characters, "
            "preventing context stuffing attacks.",
        ),
        (
            "Session turn limit:",
            "Sessions expire after 30 turns with full " "sensitive data cleanup.",
        ),
        (
            "Luhn check + CVV/expiry validation:",
            "Card numbers are validated via "
            "Luhn algorithm, CVV length is enforced (3 for standard, 4 for Amex), "
            "and expired cards are rejected.",
        ),
    ]
    for label, text in val_items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(label + " ")
        run.bold = True
        p.add_run(text)

    doc.add_heading("Session Security", level=2)
    sess_items = [
        (
            "Sensitive data cleanup on close:",
            "_close_session() clears account_data, "
            "card_details, and collected_name from working memory when the session "
            "transitions to CLOSED -- whether from successful payment, lockout, "
            "decline, or turn limit.",
        ),
        (
            "Shared verification counter:",
            "A single counter tracks failures across "
            "name AND secondary factor verification (max 3 total). Prevents "
            "brute-force enumeration.",
        ),
        (
            "Forced tool use:",
            'tool_choice: "any" ensures Claude always calls the '
            "extraction tool when tools are available, preventing the LLM from "
            "bypassing entity extraction and generating uncontrolled freeform responses.",
        ),
    ]
    for label, text in sess_items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(label + " ")
        run.bold = True
        p.add_run(text)

    # --- LLM Integration ---
    doc.add_heading("LLM Integration", level=1)

    doc.add_heading("Entity Extraction via Tool Use", level=2)
    doc.add_paragraph("Claude extracts entities through two tools with strict schemas:")
    doc.add_paragraph(
        "extract_entities -- account_id, name, DOB, Aadhaar, pincode, payment "
        "amount/intent, decline/affirmative signals",
        style="List Bullet",
    )
    doc.add_paragraph(
        "extract_card_details -- cardholder name, card number, CVV, expiry",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Progressive tool disclosure: only tools valid for the current state are "
        "exposed. Card collection state gets extract_card_details; terminal "
        "states get no tools."
    )

    doc.add_heading("State-Specific System Prompts", level=2)
    doc.add_paragraph(
        "Each ConversationState maps to a focused system prompt that tells Claude "
        "exactly what to do in that state. This prevents hallucination of "
        "irrelevant actions."
    )

    doc.add_heading("Deterministic Fallback Chain", level=2)
    doc.add_paragraph(
        "If the Claude API fails (rate limit, network error, server error):"
    )
    doc.add_paragraph(
        "Retry with exponential backoff + jitter (up to 3 attempts)",
        style="List Number",
    )
    doc.add_paragraph(
        "Fall back to template responses (one per state)", style="List Number"
    )
    doc.add_paragraph(
        "The conversation continues -- never crashes", style="List Number"
    )

    # --- Key Decisions ---
    doc.add_heading("Key Decisions", level=1)

    doc.add_heading("1. Case-Sensitive Exact Name Matching", level=2)
    doc.add_paragraph(
        "provided_name == account_data.full_name with no normalization. The "
        "assignment requires strict matching with no fuzzy workarounds."
    )

    doc.add_heading("2. Shared Verification Counter", level=2)
    doc.add_paragraph(
        "A single counter tracks failures across name AND secondary factor "
        "verification (max 3 total). Separate counters would allow 6 attempts "
        "-- too lenient for security."
    )

    doc.add_heading("3. Dependency Injection for Both API and LLM", level=2)
    doc.add_paragraph(
        "PaymentAPIClientBase and LLMClientBase are abstract. MockLLMClient "
        "returns pre-configured LLMResponse objects for deterministic unit "
        "testing. No LLM calls in unit tests."
    )

    doc.add_heading("4. Incremental Card Detail Collection", level=2)
    doc.add_paragraph(
        "Card fields are merged across messages. Users can provide all details "
        "at once or one at a time."
    )

    # --- Testing ---
    doc.add_heading("Testing Strategy (3 Layers)", level=1)
    add_table(
        doc,
        ["Layer", "Location", "What", "Speed"],
        [
            [
                "Unit",
                "tests/unit/",
                "State transitions with MockLLMClient, memory, PII scanner",
                "<0.3s",
            ],
            [
                "Integration",
                "tests/integration/",
                "Full flows with real Claude + real Prodigal API",
                "~30s",
            ],
            [
                "Eval",
                "eval/",
                "LLM-as-judge behavioral scoring across scenarios",
                "~60s",
            ],
        ],
    )

    # --- Tradeoffs ---
    doc.add_heading("Tradeoffs Accepted", level=1)
    tradeoffs = [
        "Synchronous API calls -- acceptable for CLI, not for concurrent web server use.",
        "Two LLM calls per turn -- tool_use response + follow-up text response. Could be optimized to single call with prompt engineering.",
        "No conversation summarization -- overflow window exists but summary generation isn't triggered automatically (would require an additional LLM call).",
    ]
    for i, t in enumerate(tradeoffs, 1):
        doc.add_paragraph(f"{i}. {t}")

    # --- Future Improvements ---
    doc.add_heading("What I Would Improve With More Time", level=1)
    improvements = [
        "Async support -- make both API client and LLM client async for web server use.",
        "Streaming responses -- use Claude's streaming API for better UX.",
        "Automatic conversation summarization -- trigger when overflow exceeds threshold.",
        "Rate limiting -- cooldown between verification attempts.",
        "Card tokenization -- never store raw card numbers, even temporarily; use a tokenization service.",
        "Observability -- structured logging, OpenTelemetry traces.",
        "Multi-language support -- i18n for user-facing messages.",
    ]
    for i, item in enumerate(improvements, 1):
        doc.add_paragraph(f"{i}. {item}")

    doc.save("DESIGN.docx")
    print("DESIGN.docx generated successfully")


if __name__ == "__main__":
    build_docx()
